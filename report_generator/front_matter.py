import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_generator.styles import (
    COLOR_NAVY, COLOR_SLATE, COLOR_BODY, COLOR_MUTED,
    HEX_PRIMARY, HEX_LIGHT_BG, set_cell_background, set_cell_margins, set_cell_borders
)

def add_front_matter(doc):
    """Generates all standard university B.Tech front matter pages."""
    
    # --------------------------------------------------------------------------
    # 1. TITLE PAGE
    # --------------------------------------------------------------------------
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(24)
    p_top.paragraph_format.space_after = Pt(8)
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_top.add_run("FINAL YEAR B.TECH MAJOR PROJECT REPORT")
    r_uni.font.name = "Arial"
    r_uni.font.size = Pt(13)
    r_uni.font.bold = True
    r_uni.font.color.rgb = COLOR_SLATE

    # Title Box Table
    t_title = doc.add_table(rows=1, cols=1)
    t_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_title = t_title.cell(0, 0)
    c_title.width = Inches(6.25)
    set_cell_background(c_title, HEX_LIGHT_BG)
    set_cell_margins(c_title, top=200, bottom=200, left=200, right=200)
    set_cell_borders(c_title, top="single", bottom="single", left="single", right="single",
                     color="1B365D", sz="16")
    
    p_main = c_title.paragraphs[0]
    p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_main.paragraph_format.space_before = Pt(6)
    p_main.paragraph_format.space_after = Pt(8)
    p_main.paragraph_format.line_spacing = 1.25
    r_title = p_main.add_run("HEARTCARE AI: AN INTEGRATED MACHINE LEARNING AND GEOSPATIAL INTELLIGENCE PLATFORM FOR CARDIOVASCULAR RISK STRATIFICATION AND CLINICAL DECISION SUPPORT")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = c_title.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run("Multi-Biomarker LightGBM Inference Engine • Longitudinal Patient Telemetry • What-If Intervention Simulator • Geospatial Emergency Center Routing")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(10.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SLATE

    p_sub_text = doc.add_paragraph()
    p_sub_text.paragraph_format.space_before = Pt(28)
    p_sub_text.paragraph_format.space_after = Pt(12)
    p_sub_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_req = p_sub_text.add_run("Submitted in partial fulfillment of the requirements for the award of the degree of\n")
    r_req.font.name = "Calibri"
    r_req.font.size = Pt(11)
    r_req.font.color.rgb = COLOR_BODY
    
    r_deg = p_sub_text.add_run("BACHELOR OF TECHNOLOGY\nIN\nCOMPUTER SCIENCE AND ENGINEERING")
    r_deg.font.name = "Arial"
    r_deg.font.size = Pt(12)
    r_deg.font.bold = True
    r_deg.font.color.rgb = COLOR_NAVY

    # Submission Metadata Table
    p_meta_space = doc.add_paragraph()
    p_meta_space.paragraph_format.space_before = Pt(24)

    t_meta = doc.add_table(rows=1, cols=2)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_left, c_right = t_meta.rows[0].cells
    c_left.width = Inches(3.1)
    c_right.width = Inches(3.1)
    for c in [c_left, c_right]:
        set_cell_margins(c, top=80, bottom=80, left=80, right=80)
        set_cell_borders(c)

    p_left = c_left.paragraphs[0]
    p_left.paragraph_format.line_spacing = 1.2
    r_by = p_left.add_run("SUBMITTED BY:\n")
    r_by.font.name = "Arial"
    r_by.font.size = Pt(10)
    r_by.font.bold = True
    r_by.font.color.rgb = COLOR_SLATE
    r_cand = p_left.add_run("Shiva Pranith Sai Vadicherla\nRoll No: [TO BE PROVIDED]\nDepartment of CSE")
    r_cand.font.name = "Calibri"
    r_cand.font.size = Pt(10.5)

    p_right = c_right.paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.line_spacing = 1.2
    r_guide = p_right.add_run("UNDER THE GUIDANCE OF:\n")
    r_guide.font.name = "Arial"
    r_guide.font.size = Pt(10)
    r_guide.font.bold = True
    r_guide.font.color.rgb = COLOR_SLATE
    r_guide_name = p_right.add_run("[Project Guide Name: TO BE PROVIDED]\nDesignation / Department of CSE\nAcademic Year: 2025 – 2026")
    r_guide_name.font.name = "Calibri"
    r_guide_name.font.size = Pt(10.5)

    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(36)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dept = p_foot.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\n[INSTITUTION NAME: TO BE PROVIDED]\nACADEMIC YEAR 2025–2026")
    r_dept.font.name = "Arial"
    r_dept.font.size = Pt(11)
    r_dept.font.bold = True
    r_dept.font.color.rgb = COLOR_NAVY

    # --------------------------------------------------------------------------
    # 2. CERTIFICATE
    # --------------------------------------------------------------------------
    doc.add_page_break()
    p_cert_head = doc.add_paragraph()
    p_cert_head.paragraph_format.space_before = Pt(16)
    p_cert_head.paragraph_format.space_after = Pt(16)
    p_cert_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cert_title = p_cert_head.add_run("CERTIFICATE OF ORIGINAL WORK")
    r_cert_title.font.name = "Arial"
    r_cert_title.font.size = Pt(16)
    r_cert_title.font.bold = True
    r_cert_title.font.color.rgb = COLOR_NAVY

    p_cert_body = doc.add_paragraph()
    p_cert_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_cert_body.paragraph_format.line_spacing = 1.25
    p_cert_body.paragraph_format.space_after = Pt(14)
    r_c1 = p_cert_body.add_run(
        "This is to certify that the project report entitled "
    )
    r_c1.font.name = "Calibri"
    r_c1.font.size = Pt(11)
    
    r_c_bold = p_cert_body.add_run('"HEARTCARE AI: AN INTEGRATED MACHINE LEARNING AND GEOSPATIAL INTELLIGENCE PLATFORM FOR CARDIOVASCULAR RISK STRATIFICATION AND CLINICAL DECISION SUPPORT" ')
    r_c_bold.font.name = "Calibri"
    r_c_bold.font.size = Pt(11)
    r_c_bold.font.bold = True

    r_c2 = p_cert_body.add_run(
        "submitted by Shiva Pranith Sai Vadicherla [Roll Number: TO BE PROVIDED] in partial fulfillment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering is a bonafide record of the work carried out by them under my supervision and guidance during the academic year 2025–2026.\n\n"
        "To the best of our knowledge, the results and architectural implementations embodied in this report have not been submitted to any other university or institute for the award of any degree, diploma, or certificate."
    )
    r_c2.font.name = "Calibri"
    r_c2.font.size = Pt(11)

    p_sig_space = doc.add_paragraph()
    p_sig_space.paragraph_format.space_before = Pt(40)

    t_sig = doc.add_table(rows=2, cols=2)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in t_sig.rows:
        for c in row.cells:
            c.width = Inches(3.1)
            set_cell_margins(c, top=80, bottom=80, left=80, right=80)
            set_cell_borders(c)

    p_s1 = t_sig.rows[0].cells[0].paragraphs[0]
    p_s1.add_run("______________________________\nSignature of Project Guide\n[Name: TO BE PROVIDED]\nDesignation, Dept of CSE").font.size = Pt(10)

    p_s2 = t_sig.rows[0].cells[1].paragraphs[0]
    p_s2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_s2.add_run("______________________________\nSignature of Head of Department\n[Name: TO BE PROVIDED]\nProfessor & Head, Dept of CSE").font.size = Pt(10)

    p_s3 = t_sig.rows[1].cells[0].paragraphs[0]
    p_s3.paragraph_format.space_before = Pt(30)
    p_s3.add_run("______________________________\nSignature of Internal Examiner\nDate:").font.size = Pt(10)

    p_s4 = t_sig.rows[1].cells[1].paragraphs[0]
    p_s4.paragraph_format.space_before = Pt(30)
    p_s4.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_s4.add_run("______________________________\nSignature of External Examiner\nDate:").font.size = Pt(10)

    # --------------------------------------------------------------------------
    # 3. DECLARATION
    # --------------------------------------------------------------------------
    doc.add_page_break()
    p_dec_head = doc.add_paragraph()
    p_dec_head.paragraph_format.space_before = Pt(16)
    p_dec_head.paragraph_format.space_after = Pt(16)
    p_dec_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_dec_title = p_dec_head.add_run("STUDENT DECLARATION")
    r_dec_title.font.name = "Arial"
    r_dec_title.font.size = Pt(16)
    r_dec_title.font.bold = True
    r_dec_title.font.color.rgb = COLOR_NAVY

    p_dec = doc.add_paragraph()
    p_dec.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_dec.paragraph_format.line_spacing = 1.25
    r_d_text = p_dec.add_run(
        "I, Shiva Pranith Sai Vadicherla, student of Bachelor of Technology in Computer Science and Engineering, hereby declare that the project entitled "
        '"HeartCare AI: An Integrated Machine Learning and Geospatial Intelligence Platform for Cardiovascular Risk Stratification and Clinical Decision Support" '
        "is an authentic record of our own research and software engineering work carried out under the supervision of my Project Guide.\n\n"
        "I affirm that this report is composed in original phrasing and adheres to established academic integrity and non-plagiarism standards. "
        "All data sources, machine learning algorithms (LightGBM, Scikit-Learn), clinical heuristic datasets (Cleveland & UCI Heart Disease benchmarks), "
        "and architectural libraries have been duly acknowledged and cited according to the Harvard referencing system.\n\n"
        "I further affirm that this platform is engineered strictly as an assistive clinical decision support tool and is not intended to replace licensed medical practitioner diagnosis or laboratory verification."
    )
    r_d_text.font.name = "Calibri"
    r_d_text.font.size = Pt(11)

    p_d_sign = doc.add_paragraph()
    p_d_sign.paragraph_format.space_before = Pt(45)
    r_ds = p_d_sign.add_run("Date: 23 September 2026\nPlace: [City / Campus: TO BE PROVIDED]\n\n\n_____________________________________\nShiva Pranith Sai Vadicherla\nRoll Number: [TO BE PROVIDED]")
    r_ds.font.name = "Calibri"
    r_ds.font.size = Pt(11)
    r_ds.font.bold = True

    # --------------------------------------------------------------------------
    # 4. ACKNOWLEDGEMENT
    # --------------------------------------------------------------------------
    doc.add_page_break()
    p_ack_head = doc.add_paragraph()
    p_ack_head.paragraph_format.space_before = Pt(16)
    p_ack_head.paragraph_format.space_after = Pt(16)
    p_ack_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ack_title = p_ack_head.add_run("ACKNOWLEDGEMENT")
    r_ack_title.font.name = "Arial"
    r_ack_title.font.size = Pt(16)
    r_ack_title.font.bold = True
    r_ack_title.font.color.rgb = COLOR_NAVY

    p_ack = doc.add_paragraph()
    p_ack.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ack.paragraph_format.line_spacing = 1.25
    r_ack = p_ack.add_run(
        "The completion of this final-year major engineering project has been a deeply enriching milestone, made possible through the generous intellectual guidance, technical resources, and encouragement of numerous individuals.\n\n"
        "I express my profound gratitude to my Project Guide, [Project Guide Name: TO BE PROVIDED], Department of Computer Science and Engineering, for their invaluable mentorship, insightful critique, and sustained guidance throughout the formulation, machine learning model calibration, and cloud full-stack architecture design of HeartCare AI.\n\n"
        "I extend my sincere thanks to Professor & Head of Department, [HOD Name: TO BE PROVIDED], for fostering an intellectually rigorous research environment and providing the necessary departmental laboratory computing infrastructure.\n\n"
        "I also thank the esteemed Principal / Director and the faculty members of the Department of Computer Science and Engineering for their direct and indirect contributions, encouragement, and constructive evaluations during project review stages.\n\n"
        "Lastly, I express our heartfelt appreciation to my family and peers for their continuous moral support and understanding throughout the course of this academic degree."
    )
    r_ack.font.name = "Calibri"
    r_ack.font.size = Pt(11)

    # --------------------------------------------------------------------------
    # 5. ABSTRACT & KEYWORDS
    # --------------------------------------------------------------------------
    doc.add_page_break()
    p_abs_head = doc.add_paragraph()
    p_abs_head.paragraph_format.space_before = Pt(16)
    p_abs_head.paragraph_format.space_after = Pt(16)
    p_abs_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_abs_title = p_abs_head.add_run("ABSTRACT")
    r_abs_title.font.name = "Arial"
    r_abs_title.font.size = Pt(16)
    r_abs_title.font.bold = True
    r_abs_title.font.color.rgb = COLOR_NAVY

    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.paragraph_format.line_spacing = 1.25
    r_abs = p_abs.add_run(
        "Cardiovascular diseases (CVDs), predominantly heart failure and coronary artery disease, represent the leading cause of global mortality, accounting for approximately 17.9 million deaths annually. Clinical decision-making in cardiology often grapples with complex, multi-factorial biomarker interactions where conventional static risk scoring tables fail to capture non-linear prognostic correlations or provide dynamic, actionable patient guidance. Furthermore, existing digital health portals rarely synthesize real-time clinical predictive modeling with explainable AI feature attribution, interactive counterfactual ('What-If') simulations, and live geospatial emergency hospital routing into a unified, secure system.\n\n"
        "To bridge this critical translational gap, this project develops HeartCare AI, an evidence-based clinical intelligence platform designed for multi-biomarker cardiovascular risk prediction, longitudinal patient telemetry, and emergency triage routing. The system is architected around a calibrated LightGBM (Light Gradient Boosting Machine) multi-class classification engine trained on 13 standardized clinical features derived from the benchmark Cleveland and UCI Heart Disease cohorts, augmented with cardiorenal biomarkers including left ventricular ejection fraction and serum creatinine. An automated fallback engine operationalizing Framingham and ACC/AHA cardiovascular guidelines ensures seamless high-availability inference even during isolated container restarts.\n\n"
        "On empirical evaluation across 297 clean clinical benchmark patient records, the integrated LightGBM classification engine achieves an overall diagnostic Accuracy of 81.82%, Precision of 76.77%, Sensitivity (Recall) of 86.86%, Specificity of 77.50%, F1-Score of 81.51%, and an Area Under the ROC Curve (ROC-AUC) of 91.70%. The platform's high recall is of paramount clinical importance, effectively minimizing hazardous false-negative diagnostic oversights. Complementing the ML engine, the full-stack architecture comprises a high-throughput FastAPI (Python 3.11/3.14) asynchronous backend, a thread-safe connection-pooled MongoDB Atlas persistence layer enforcing strict cross-user data isolation, and a glassmorphic React 19 single-page frontend featuring an HTML5 canvas cardiac rhythm visualizer and interactive What-If parameter simulators. Live geospatial triage is powered by OpenStreetMap Overpass and Nominatim APIs computing geodesic Haversine distance and travel ETAs to verified cardiology super-speciality centers across India.\n\n"
        "Rigorous end-to-end quality assurance auditing across 104 executed test cases confirmed 103 passed tests with sub-second execution latencies (842ms serial, 559ms concurrent). The developed platform provides an extensible, production-ready blueprint for next-generation clinical decision support systems."
    )
    r_abs.font.name = "Calibri"
    r_abs.font.size = Pt(11)

    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(14)
    p_kw.paragraph_format.space_after = Pt(8)
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.name = "Arial"
    r_kwh.font.size = Pt(11)
    r_kwh.font.bold = True
    r_kwh.font.color.rgb = COLOR_NAVY

    r_kw = p_kw.add_run(
        "Cardiovascular Risk Stratification, Machine Learning, LightGBM Classifier, Explainable AI, Clinical Decision Support Systems, FastAPI, React 19, MongoDB Atlas, Geodesic Haversine Formula, Emergency Medical Triage."
    )
    r_kw.font.name = "Calibri"
    r_kw.font.size = Pt(10.5)
    r_kw.font.italic = True

    # --------------------------------------------------------------------------
    # 6. TABLE OF CONTENTS
    # --------------------------------------------------------------------------
    doc.add_page_break()
    p_toc_head = doc.add_paragraph()
    p_toc_head.paragraph_format.space_before = Pt(16)
    p_toc_head.paragraph_format.space_after = Pt(14)
    p_toc_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_toc_title = p_toc_head.add_run("TABLE OF CONTENTS")
    r_toc_title.font.name = "Arial"
    r_toc_title.font.size = Pt(16)
    r_toc_title.font.bold = True
    r_toc_title.font.color.rgb = COLOR_NAVY

    toc_items = [
        ("Certificate of Original Work", "ii"),
        ("Student Declaration", "iii"),
        ("Acknowledgement", "iv"),
        ("Abstract", "v"),
        ("List of Figures", "viii"),
        ("List of Tables", "ix"),
        ("CHAPTER 1: INTRODUCTION", "1"),
        ("    1.1 Introduction", "1"),
        ("    1.2 Background of the Project", "2"),
        ("    1.3 Problem Statement", "3"),
        ("    1.4 Motivation", "4"),
        ("    1.5 Objectives", "5"),
        ("    1.6 Scope of the Project", "6"),
        ("    1.7 Significance of the Project", "7"),
        ("    1.8 Limitations", "8"),
        ("    1.9 Organization of the Report", "9"),
        ("CHAPTER 2: LITERATURE SURVEY", "10"),
        ("    2.1 Introduction", "10"),
        ("    2.2 Review of Existing Research Papers", "11"),
        ("    2.3 Existing Systems", "14"),
        ("    2.4 Comparison of Existing Systems", "16"),
        ("    2.5 Research Gap", "18"),
        ("    2.6 Proposed System", "19"),
        ("    2.7 Summary", "21"),
        ("CHAPTER 3: PROJECT FLOW AND METHODOLOGY", "22"),
        ("    3.1 Introduction", "22"),
        ("    3.2 Overall Project Workflow", "23"),
        ("    3.3 Proposed Methodology", "25"),
        ("    3.4 Data Collection", "26"),
        ("    3.5 Data Preprocessing", "28"),
        ("    3.6 Feature Engineering", "30"),
        ("    3.7 Machine Learning Methodology", "32"),
        ("    3.8 Model Training and Evaluation", "35"),
        ("    3.9 Explainable AI Methodology", "38"),
        ("    3.10 Backend and Frontend Workflow", "40"),
        ("    3.11 Database Workflow", "42"),
        ("    3.12 Overall System Flow", "44"),
        ("CHAPTER 4: SYSTEM DESIGN", "46"),
        ("    4.1 Introduction", "46"),
        ("    4.2 System Architecture", "47"),
        ("    4.3 Architecture Components", "49"),
        ("    4.4 Functional Requirements", "52"),
        ("    4.5 Non-Functional Requirements", "54"),
        ("    4.6 Use Case Diagram and Specifications", "56"),
        ("    4.7 Data Flow Diagrams (DFD Level 0 and Level 1)", "59"),
        ("    4.8 System Flowchart", "62"),
        ("    4.9 Database Design and Entity-Relationship (ER) Schema", "65"),
        ("    4.10 Module Design", "68"),
        ("    4.11 API Design and Contracts", "71"),
        ("    4.12 Security Architecture and Threat Mitigation", "74"),
        ("CHAPTER 5: IMPLEMENTATION", "77"),
        ("    5.1 Introduction", "77"),
        ("    5.2 Development Environment", "78"),
        ("    5.3 Technology Stack", "80"),
        ("    5.4 Frontend Implementation", "83"),
        ("    5.5 Backend Implementation", "87"),
        ("    5.6 Database Implementation", "91"),
        ("    5.7 Machine Learning Implementation", "95"),
        ("    5.8 Model Prediction Process", "99"),
        ("    5.9 Explainable AI Implementation", "102"),
        ("    5.10 Feature Implementation", "105"),
        ("    5.11 API Integration", "108"),
        ("    5.12 Frontend-Backend Integration", "111"),
        ("    5.13 Implementation Summary", "114"),
        ("CHAPTER 6: RESULT ANALYSIS", "116"),
        ("    6.1 Introduction", "116"),
        ("    6.2 Dataset Summary", "117"),
        ("    6.3 Experimental Setup", "119"),
        ("    6.4 Model Evaluation Metrics", "121"),
        ("    6.5 Accuracy, Precision, Recall, and F1-Score", "123"),
        ("    6.6 Confusion Matrix Analysis", "125"),
        ("    6.7 Model Comparison Benchmark", "127"),
        ("    6.8 Explainability Results", "129"),
        ("    6.9 Sample Prediction Clinical Profiles", "131"),
        ("    6.10 System Performance and Latency Benchmarks", "134"),
        ("    6.11 Discussion", "136"),
        ("    6.12 Limitations", "138"),
        ("CHAPTER 7: TESTING", "140"),
        ("    7.1 Introduction", "140"),
        ("    7.2 Testing Strategy", "141"),
        ("    7.3 Unit Testing", "143"),
        ("    7.4 Integration Testing", "145"),
        ("    7.5 System Testing", "147"),
        ("    7.6 UI Testing", "149"),
        ("    7.7 API Testing", "151"),
        ("    7.8 Functional Testing", "153"),
        ("    7.9 Non-Functional Testing", "155"),
        ("    7.10 Test Case Tables", "157"),
        ("    7.11 Test Results and Defect Analysis", "162"),
        ("    7.12 Testing Summary", "165"),
        ("CHAPTER 8: MAINTENANCE OF THE SYSTEM", "167"),
        ("    8.1 Introduction", "167"),
        ("    8.2 Types of Maintenance", "168"),
        ("    8.3 Corrective Maintenance", "170"),
        ("    8.4 Adaptive Maintenance", "172"),
        ("    8.5 Perfective Maintenance", "174"),
        ("    8.6 Preventive Maintenance", "176"),
        ("    8.7 Database Maintenance", "178"),
        ("    8.8 Security Maintenance", "180"),
        ("    8.9 Machine Learning Model Maintenance", "182"),
        ("    8.10 Backup and Disaster Recovery", "184"),
        ("    8.11 Future Maintenance Requirements", "186"),
        ("CHAPTER 9: DEPLOYMENT", "188"),
        ("    9.1 Introduction", "188"),
        ("    9.2 Deployment Architecture", "189"),
        ("    9.3 Deployment Requirements", "191"),
        ("    9.4 Frontend Deployment (Vercel SPA)", "193"),
        ("    9.5 Backend Deployment (Render Container)", "195"),
        ("    9.6 Database Deployment (MongoDB Atlas)", "197"),
        ("    9.7 Environment Configuration", "199"),
        ("    9.8 API Configuration and CORS", "201"),
        ("    9.9 Security During Deployment", "203"),
        ("    9.10 Deployment Verification", "205"),
        ("    9.11 Deployment Challenges and Mitigations", "207"),
        ("    9.12 Post-Deployment Maintenance", "209"),
        ("REFERENCES", "211")
    ]

    t_toc = doc.add_table(rows=len(toc_items), cols=2)
    t_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (title, page) in enumerate(toc_items):
        r_cells = t_toc.rows[idx].cells
        r_cells[0].width = Inches(5.4)
        r_cells[1].width = Inches(0.85)
        for c in r_cells:
            set_cell_margins(c, top=40, bottom=40, left=40, right=40)
            set_cell_borders(c)
        
        p0 = r_cells[0].paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.line_spacing = 1.15
        run0 = p0.add_run(title)
        run0.font.name = "Calibri"
        run0.font.size = Pt(10)
        if "CHAPTER" in title or "REFERENCES" in title or "Abstract" in title:
            run0.font.bold = True
            run0.font.color.rgb = COLOR_NAVY
        else:
            run0.font.color.rgb = COLOR_BODY

        p1 = r_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        run1 = p1.add_run(page)
        run1.font.name = "Calibri"
        run1.font.size = Pt(10)
        run1.font.color.rgb = COLOR_MUTED

    # --------------------------------------------------------------------------
    # 7. LIST OF FIGURES & LIST OF TABLES
    # --------------------------------------------------------------------------
    doc.add_page_break()
    p_lof_head = doc.add_paragraph()
    p_lof_head.paragraph_format.space_before = Pt(16)
    p_lof_head.paragraph_format.space_after = Pt(12)
    p_lof_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_lof_title = p_lof_head.add_run("LIST OF FIGURES")
    r_lof_title.font.name = "Arial"
    r_lof_title.font.size = Pt(16)
    r_lof_title.font.bold = True
    r_lof_title.font.color.rgb = COLOR_NAVY

    figures_list = [
        ("Figure 3.1", "HeartCare AI End-to-End High-Level System Workflow", "24"),
        ("Figure 3.2", "LightGBM Leaf-Wise Tree Growth Strategy vs Level-Wise Splitting", "34"),
        ("Figure 3.3", "Clinical Telemetry and Biomarker Processing Pipeline", "41"),
        ("Figure 4.1", "Multi-Tier Cloud Microservices System Architecture Diagram", "48"),
        ("Figure 4.2", "System Use Case Diagram for Clinicians and Patients", "57"),
        ("Figure 4.3", "Data Flow Diagram (DFD Level 0 Context Diagram)", "60"),
        ("Figure 4.4", "Data Flow Diagram (DFD Level 1 Detailed Functional Flow)", "61"),
        ("Figure 4.5", "End-to-End Clinical Diagnostic System Flowchart", "63"),
        ("Figure 4.6", "MongoDB Atlas Entity-Relationship (ER) Schema Design", "66"),
        ("Figure 6.1", "ROC-AUC Curve for LightGBM Binary Classification (91.70% AUC)", "124"),
        ("Figure 6.2", "Binary Classification Confusion Matrix Heatmap", "126"),
        ("Figure 6.3", "5-Class Disease Severity Confusion Matrix (Stages 0 to 4)", "127"),
        ("Figure 6.4", "SHAP Biomarker Attribution and Feature Impact Breakdown", "130"),
        ("Figure 7.1", "V-Model Quality Assurance and Testing Hierarchy", "142"),
        ("Figure 9.1", "Multi-Cloud Containerized Deployment Architecture (Vercel + Render + Atlas)", "190")
    ]

    t_lof = doc.add_table(rows=len(figures_list), cols=2)
    t_lof.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (f_num, f_desc, f_page) in enumerate(figures_list):
        r_cells = t_lof.rows[idx].cells
        r_cells[0].width = Inches(5.4)
        r_cells[1].width = Inches(0.85)
        for c in r_cells:
            set_cell_margins(c, top=40, bottom=40, left=40, right=40)
            set_cell_borders(c)
        
        p0 = r_cells[0].paragraphs[0]
        r_fnum = p0.add_run(f"{f_num}: ")
        r_fnum.font.name = "Calibri"
        r_fnum.font.size = Pt(10)
        r_fnum.font.bold = True
        r_fnum.font.color.rgb = COLOR_NAVY
        r_fdesc = p0.add_run(f_desc)
        r_fdesc.font.name = "Calibri"
        r_fdesc.font.size = Pt(10)
        r_fdesc.font.color.rgb = COLOR_BODY

        p1 = r_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_fpage = p1.add_run(f_page)
        r_fpage.font.name = "Calibri"
        r_fpage.font.size = Pt(10)
        r_fpage.font.color.rgb = COLOR_MUTED

    p_lot_space = doc.add_paragraph()
    p_lot_space.paragraph_format.space_before = Pt(20)

    p_lot_head = doc.add_paragraph()
    p_lot_head.paragraph_format.space_before = Pt(12)
    p_lot_head.paragraph_format.space_after = Pt(12)
    p_lot_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_lot_title = p_lot_head.add_run("LIST OF TABLES")
    r_lot_title.font.name = "Arial"
    r_lot_title.font.size = Pt(16)
    r_lot_title.font.bold = True
    r_lot_title.font.color.rgb = COLOR_NAVY

    tables_list = [
        ("Table 2.1", "Comparison Matrix of Existing Cardiovascular Prediction Systems", "17"),
        ("Table 3.1", "Description of 13 Cleveland Clinical Features and Encodings", "27"),
        ("Table 3.2", "Extended Cardiorenal Telemetry and Lifestyle Biomarkers", "29"),
        ("Table 4.1", "Functional Requirements Specification (FR1 to FR12)", "53"),
        ("Table 4.2", "Non-Functional Requirements Specification (NFR1 to NFR8)", "55"),
        ("Table 4.3", "RESTful API Endpoints and Schema Definitions", "72"),
        ("Table 5.1", "Complete Development and Deployment Technology Stack", "81"),
        ("Table 6.1", "UCI Cleveland Dataset Feature Statistics Summary", "118"),
        ("Table 6.2", "Primary Clinical Diagnosis Performance Metrics (Healthy vs Disease)", "123"),
        ("Table 6.3", "Multiclass Severity Classification Report (Stages 0 to 4)", "126"),
        ("Table 6.4", "Benchmarking Algorithm Performance on Cleveland Dataset", "128"),
        ("Table 6.5", "Performance and Latency Benchmarks under Concurrent Loads", "135"),
        ("Table 7.1", "End-to-End QA Audit Test Execution Summary Statistics (104 Tests)", "142"),
        ("Table 7.2", "Detailed Functional and Security Test Case Log (TC-01 to TC-15)", "158")
    ]

    t_lot = doc.add_table(rows=len(tables_list), cols=2)
    t_lot.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (t_num, t_desc, t_page) in enumerate(tables_list):
        r_cells = t_lot.rows[idx].cells
        r_cells[0].width = Inches(5.4)
        r_cells[1].width = Inches(0.85)
        for c in r_cells:
            set_cell_margins(c, top=40, bottom=40, left=40, right=40)
            set_cell_borders(c)
        
        p0 = r_cells[0].paragraphs[0]
        r_tnum = p0.add_run(f"{t_num}: ")
        r_tnum.font.name = "Calibri"
        r_tnum.font.size = Pt(10)
        r_tnum.font.bold = True
        r_tnum.font.color.rgb = COLOR_NAVY
        r_tdesc = p0.add_run(t_desc)
        r_tdesc.font.name = "Calibri"
        r_tdesc.font.size = Pt(10)
        r_tdesc.font.color.rgb = COLOR_BODY

        p1 = r_cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_tpage = p1.add_run(t_page)
        r_tpage.font.name = "Calibri"
        r_tpage.font.size = Pt(10)
        r_tpage.font.color.rgb = COLOR_MUTED
