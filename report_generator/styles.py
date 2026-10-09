import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Palette
COLOR_NAVY = RGBColor(27, 54, 93)       # #1B365D - Primary headings
COLOR_SLATE = RGBColor(43, 84, 126)     # #2B547E - Subheadings
COLOR_BODY = RGBColor(40, 40, 40)       # #282828 - Body text
COLOR_MUTED = RGBColor(100, 100, 100)   # #646464 - Captions / Footers
COLOR_CRIMSON = RGBColor(185, 28, 28)   # #B91C1C - Critical / Warnings
HEX_PRIMARY = "1B365D"
HEX_SLATE = "2B547E"
HEX_LIGHT_BG = "F4F6F9"
HEX_ALT_ROW = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BORDER = "1B365D"

def set_cell_background(cell, hex_color):
    """Sets cell background fill color using XML."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Sets cell internal padding (in twips: 20 twips = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="4"):
    """Sets individual borders on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for b_name, b_val in borders.items():
        if b_val != "none":
            b_el = OxmlElement(f'w:{b_name}')
            b_el.set(qn('w:val'), b_val)
            b_el.set(qn('w:sz'), sz)
            b_el.set(qn('w:space'), '0')
            b_el.set(qn('w:color'), color)
            tcBorders.append(b_el)
        else:
            b_el = OxmlElement(f'w:{b_name}')
            b_el.set(qn('w:val'), 'none')
            tcBorders.append(b_el)
    tcPr.append(tcBorders)

def configure_document(doc):
    """Configures global margins, styles, header and footer."""
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)   # Extra 0.25" for binding gutter
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

        # Header setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("HeartCare AI: Cardiovascular Risk Stratification & Clinical Decision Platform")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED
        hrun.italic = True

        # Footer setup
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun_left = fp.add_run("Department of Computer Science & Engineering | Final-Year B.Tech Project Report\t")
        frun_left.font.name = "Calibri"
        frun_left.font.size = Pt(8.5)
        frun_left.font.color.rgb = COLOR_MUTED
        
        frun_right = fp.add_run("HeartCare.AI")
        frun_right.font.name = "Calibri"
        frun_right.font.size = Pt(8.5)
        frun_right.font.color.rgb = COLOR_NAVY
        frun_right.bold = True

def add_chapter_title(doc, chapter_num, title):
    """Adds a major Chapter Header with page break."""
    doc.add_page_break()
    p_num = doc.add_paragraph()
    p_num.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_num.paragraph_format.space_before = Pt(12)
    p_num.paragraph_format.space_after = Pt(2)
    run_num = p_num.add_run(f"CHAPTER {chapter_num}".upper())
    run_num.font.name = "Arial"
    run_num.font.size = Pt(13)
    run_num.font.bold = True
    run_num.font.color.rgb = COLOR_SLATE

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(18)
    run_title = p_title.add_run(title)
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY

    # Decorative bottom accent rule
    rule_table = doc.add_table(rows=1, cols=1)
    rule_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = rule_table.cell(0, 0)
    cell.width = Inches(6.25)
    set_cell_background(cell, HEX_PRIMARY)
    set_cell_margins(cell, top=20, bottom=20, left=0, right=0)
    cell.paragraphs[0].paragraph_format.space_before = Pt(0)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(6)

def add_section_heading(doc, num_str, title):
    """Adds a Level 2 Heading (e.g., 1.1 Introduction)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{num_str} {title}")
    run.font.name = "Arial"
    run.font.size = Pt(13.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY

def add_subsection_heading(doc, num_str, title):
    """Adds a Level 3 Heading (e.g., 1.1.1 Background)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(f"{num_str} {title}")
    run.font.name = "Arial"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SLATE

def add_paragraph(doc, text):
    """Adds a justified, standard academic body paragraph."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_BODY
    return p

def add_bullet(doc, title, body):
    """Adds a structured academic bullet point with bold lead-in."""
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(title)
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_NAVY

    run_b = p.add_run(body)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(11)
    run_b.font.color.rgb = COLOR_BODY
    return p

def add_callout(doc, text, title="NOTE", border_color="1B365D", bg_color="F4F6F9"):
    """Adds an academic callout box with a thick colored left accent border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.25)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    set_cell_borders(cell, top="none", bottom="none", left="single", right="none", 
                     color=border_color, sz="24")
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15

    r_title = p.add_run(f"[{title}] ")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(10.5)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY if border_color == "1B365D" else COLOR_CRIMSON

    r_body = p.add_run(text)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = COLOR_BODY

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str, caption=None):
    """Adds a shaded monospace code listing."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.25)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single",
                     color="CBD5E1", sz="4")

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)

    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run(f"Listing: {caption}")
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED

def add_styled_table(doc, headers, data, col_widths=None, caption=None):
    """
    Creates a publication-quality styled table with colored header,
    alternating row shading, and thin borders.
    """
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(4)
        p_cap.paragraph_format.keep_with_next = True
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(10)
        r_cap.font.bold = True
        r_cap.font.color.rgb = COLOR_NAVY

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h_text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], HEX_PRIMARY)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        set_cell_borders(hdr_cells[i], top="single", bottom="single", left="single", right="single",
                         color="1B365D", sz="4")

    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_col = HEX_ALT_ROW if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = ""
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(str(cell_value))
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_BODY
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
            set_cell_borders(row_cells[col_idx], top="single", bottom="single", left="single", right="single",
                             color=HEX_BORDER, sz="4")

    # Column Widths
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(4)
    p_after.paragraph_format.space_after = Pt(6)

def add_diagram_box(doc, ascii_diagram, caption):
    """
    Renders an architectural/workflow diagram in a centered boxed container
    with monospace font and a descriptive figure caption.
    """
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.25)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=160, bottom=160, left=200, right=180)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single",
                     color="94A3B8", sz="6")

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(ascii_diagram)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(15, 23, 42)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(4)
    p_cap.paragraph_format.space_after = Pt(8)
    r_cap = p_cap.add_run(f"Figure: {caption}")
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(10)
    r_cap.font.bold = True
    r_cap.font.color.rgb = COLOR_NAVY
